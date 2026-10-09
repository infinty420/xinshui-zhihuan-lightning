"""把《项目成果说明》Markdown 转为排版规整的 Word 文档"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "项目成果说明.md"
OUT = ROOT / "项目成果说明.docx"

BRAND = RGBColor(0x1F, 0x3A, 0x66)
MUTED = RGBColor(0x5A, 0x63, 0x72)


def set_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_rich(par, text, size=10.5):
    """处理 **粗体** 与 `代码` 标记"""
    for seg in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = par.add_run(seg[2:-2])
            set_font(r, size=size, bold=True)
        elif seg.startswith("`") and seg.endswith("`"):
            r = par.add_run(seg[1:-1])
            set_font(r, name="Consolas", size=size - 0.5, color=RGBColor(0xC0, 0x39, 0x2B))
        else:
            r = par.add_run(seg)
            set_font(r, size=size)


def main() -> int:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.6)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        # 分割线
        if line == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("—" * 34)
            set_font(r, size=8, color=RGBColor(0xC5, 0xCC, 0xD8))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 表格
        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncol)
                table.style = "Light Grid Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    cells = table.add_row().cells
                    for ci in range(ncol):
                        txt = row[ci] if ci < len(row) else ""
                        cp = cells[ci].paragraphs[0]
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(2)
                        add_rich(cp, txt, size=9.5)
                        if ri == 0:
                            for r in cp.runs:
                                r.font.bold = True
                doc.add_paragraph()
            continue

        # 代码块
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            for bi, bl in enumerate(block):
                r = p.add_run(bl + ("\n" if bi < len(block) - 1 else ""))
                set_font(r, name="Consolas", size=9)
                r.font.color.rgb = RGBColor(0x33, 0x3D, 0x4D)
            continue

        # 标题
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(16)
                add_rich(p, text, size=18)
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = BRAND
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                add_rich(p, text, size=13.5)
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = BRAND
            elif level == 3:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                add_rich(p, text, size=11.5)
                for r in p.runs:
                    r.font.bold = True
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                add_rich(p, text, size=10.5)
                for r in p.runs:
                    r.font.bold = True
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(8)
            add_rich(p, line.lstrip("> ").strip(), size=9.5)
            for r in p.runs:
                r.font.color.rgb = MUTED
                r.font.italic = True
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(2))
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(3))
            i += 1
            continue

        # 空行
        if not line:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.45
        add_rich(p, line)
        i += 1

    doc.save(OUT)
    print("已生成：", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
