"""
多格式工资表适配器
==================
把各类客户形态各异的工资表（xlsx / xls / csv，长表 / 宽表 / 单月表）
解析并映射为统一 schema，供计算内核消费。

支持的三种表形态：
  1. 长表：一行 = 一名员工 × 一个月（含「月份」列）
  2. 宽表：一行 = 一名员工，月份横向铺开（如「1月基本工资」「2月应发」…）
  3. 单月表：一行 = 一名员工，只有一个月（月份由参数或文件名指定）
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .policy import D
from .schema import FIELD_MAP, INCOME_KEYS, MappingResult, build_mapping, match_field

MONTH_PATTERNS = [
    re.compile(r"(?:(\d{4})\s*[年\-/.])?\s*(\d{1,2})\s*月份?"),   # 2026年1月 / 1月 / 2026-1
    re.compile(r"(?:(\d{4}))\s*[-/]?\s*(\d{1,2})(?!\d)"),          # 2026-01 / 202601
]
MONTH_WORDS = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "oct": 10,
    "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}


@dataclass
class SheetData:
    name: str
    headers: list[str]
    rows: list[dict[str, Any]]
    header_row_index: int = 0

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.rows), len(self.headers)


@dataclass
class ParseReport:
    sheet_name: str = ""
    layout: str = ""                              # long / wide / single
    mapping: MappingResult = field(default_factory=MappingResult)
    row_count: int = 0
    records: int = 0
    notes: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------
def _clean(value) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        return v or None
    return value


def _rows_from_openpyxl(path: Path) -> dict[str, list[list[Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out: dict[str, list[list[Any]]] = {}
    for ws in wb.worksheets:
        data = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
        out[ws.title] = [r for r in data if any(v is not None for v in r)]
    wb.close()
    return out


def _rows_from_xls(path: Path) -> dict[str, list[list[Any]]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    out: dict[str, list[list[Any]]] = {}
    for sh in book.sheets():
        data = [[_clean(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
        out[sh.name] = [r for r in data if any(v is not None for v in r)]
    return out


def _rows_from_csv(path: Path) -> dict[str, list[list[Any]]]:
    for enc in ("utf-8-sig", "gbk", "gb18030", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                data = [[_clean(c) for c in row] for row in csv.reader(f)]
            return {path.stem: [r for r in data if any(v is not None for v in r)]}
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别 CSV 编码：{path}")


def load_workbook(path: str | Path) -> dict[str, list[list[Any]]]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _rows_from_openpyxl(p)
    if ext == ".xls":
        return _rows_from_xls(p)
    if ext in (".csv", ".txt"):
        return _rows_from_csv(p)
    raise ValueError(f"不支持的文件格式：{ext}（支持 xlsx / xlsm / xls / csv）")


def list_sheets(path: str | Path) -> list[str]:
    return list(load_workbook(path).keys())


# ---------------------------------------------------------------------------
# 表头探测
# ---------------------------------------------------------------------------
def detect_header_row(matrix: list[list[Any]], scan: int = 15) -> int:
    """在前 scan 行中找最像表头的一行（能匹配上最多标准字段且非空单元格多）"""
    best_idx, best_score = 0, -1
    for i, row in enumerate(matrix[:scan]):
        cells = [c for c in row if c is not None]
        if len(cells) < 2:
            continue
        hits = sum(1 for c in cells if match_field(str(c)))
        score = hits * 10 + len(cells)
        # 纯数字行不可能是表头
        numeric = sum(1 for c in cells if isinstance(c, (int, float)))
        if numeric > len(cells) / 2:
            score -= 30
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def to_sheet(name: str, matrix: list[list[Any]]) -> SheetData:
    """把二维矩阵切成 SheetData，自动探测表头行"""
    if not matrix:
        return SheetData(name=name, headers=[], rows=[])
    hr = detect_header_row(matrix)
    raw_headers = matrix[hr]
    headers: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(raw_headers):
        label = str(h).strip() if h is not None else f"列{i + 1}"
        if not label:
            label = f"列{i + 1}"
        if label in seen:                      # 重名列加后缀，避免覆盖
            seen[label] += 1
            label = f"{label}#{seen[label]}"
        else:
            seen[label] = 0
        headers.append(label)

    rows: list[dict[str, Any]] = []
    for r in matrix[hr + 1:]:
        if not any(v is not None for v in r):
            continue
        row = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}
        rows.append(row)
    return SheetData(name=name, headers=headers, rows=rows, header_row_index=hr)


def load_sheet(path: str | Path, sheet: str | None = None) -> SheetData:
    """读取文件并返回一个可用的 SheetData（自动选 sheet）"""
    book = load_workbook(path)
    if not book:
        raise ValueError("文件为空或无可解析内容")
    if sheet and sheet in book:
        name = sheet
    else:
        # 选择行数最多的 sheet，若有多个则优先名称含“工资/薪”的
        candidates = sorted(book.items(), key=lambda kv: (-len(kv[1]), kv[0]))
        preferred = [kv for kv in candidates if re.search(r"工资|薪|工资表|salary|payroll", kv[0])]
        name = (preferred or candidates)[0][0]
    return to_sheet(name, book[name])


# ---------------------------------------------------------------------------
# 月份提取（宽表展开用）
# ---------------------------------------------------------------------------
def extract_month(header: str) -> tuple[int | None, str]:
    """从表头中剥离月份信息，返回 (月份, 剩余文本)"""
    if not header:
        return None, ""
    text = str(header).strip()
    for pat in MONTH_PATTERNS:
        for m in pat.finditer(text):
            month = int(m.group(2))
            if 1 <= month <= 12:
                rest = (text[:m.start()] + text[m.end():]).strip()
                return month, rest
    low = re.sub(r"[^a-z]", "", text.lower())
    if low in MONTH_WORDS:
        return MONTH_WORDS[low], ""
    return None, text


def detect_layout(headers: list[str], mapping: MappingResult) -> str:
    """判断表形态：long（有月份列）/ wide（月份在列名里）/ single"""
    if "month" in mapping.mapping:
        return "long"
    monthish = sum(1 for h in headers if extract_month(str(h))[0] is not None)
    if monthish >= 2:
        return "wide"
    return "single"


# ---------------------------------------------------------------------------
# 标准化
# ---------------------------------------------------------------------------
@dataclass
class StandardizeContext:
    year: int | None = None
    month: int | None = None            # 单月表 / 缺省月份
    city: str | None = None
    bonus_method: str = "auto"
    declaration_method: str = "monthly"
    defaults: dict[str, Any] = field(default_factory=dict)


def _num(value) -> Any:
    """把单元格值转成数字或 None"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).strip().replace(",", "").replace("￥", "").replace("¥", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def standardize(
    sheet: SheetData,
    ctx: StandardizeContext,
    mapping: MappingResult | None = None,
) -> tuple[list[dict[str, Any]], ParseReport]:
    """把 SheetData 转成标准化记录列表"""
    mapping = mapping or build_mapping(sheet.headers)
    layout = detect_layout(sheet.headers, mapping)
    report = ParseReport(sheet_name=sheet.name, layout=layout, mapping=mapping, row_count=len(sheet.rows))

    col = mapping.mapping                       # 标准字段 -> 原始表头
    records: list[dict[str, Any]] = []

    def make_record(row: dict[str, Any], col_map: dict[str, str]) -> dict[str, Any]:
        """按给定的列映射从一行中抽取标准字段"""
        rec: dict[str, Any] = {}
        for key, header in col_map.items():
            raw = row.get(header)
            if key in ("name", "id_no", "department", "city", "employee_id"):
                rec[key] = str(raw).strip() if raw is not None else None
            else:
                rec[key] = _num(raw)
        # 员工编号 / 身份证可能带 .0 后缀（Excel 被识别为数字）
        if rec.get("employee_id"):
            rec["employee_id"] = re.sub(r"\.0$", "", str(rec["employee_id"])).strip()
        if rec.get("id_no"):
            rec["id_no"] = re.sub(r"\.0$", "", str(rec["id_no"])).strip()
        for k, v in ctx.defaults.items():
            if rec.get(k) in (None, ""):
                rec[k] = v
        if rec.get("city") in (None, "") and ctx.city:
            rec["city"] = ctx.city
        if rec.get("year") in (None, "") and ctx.year:
            rec["year"] = ctx.year
        rec["annual_bonus_method"] = ctx.bonus_method
        rec["declaration_method"] = ctx.declaration_method
        return rec

    def base_record(row: dict[str, Any]) -> dict[str, Any]:
        return make_record(row, col)

    if layout == "long":
        for row in sheet.rows:
            rec = base_record(row)
            month = rec.get("month")
            if month in (None, ""):
                continue
            rec["month"] = int(float(month))
            records.append(rec)

    elif layout == "wide":
        # 建立：月份 -> {标准字段: 原始表头}
        month_cols: dict[int, dict[str, str]] = {}
        employee_cols: dict[str, str] = {}
        for h in sheet.headers:
            month, rest = extract_month(str(h))
            key = match_field(rest) or match_field(str(h))
            if month is not None and key:
                month_cols.setdefault(month, {})[key] = h
            elif month is None and key and key not in ("month",):
                employee_cols.setdefault(key, h)

        if not month_cols:
            report.notes.append("检测到月份列但未识别出可映射字段，已回退为单月表处理")
            for row in sheet.rows:
                rec = base_record(row)
                rec["month"] = ctx.month or 1
                records.append(rec)
        else:
            report.notes.append(f"宽表已展开为 {len(month_cols)} 个月份：{sorted(month_cols)}")
            identity_keys = ("employee_id", "name", "id_no", "department", "city", "year",
                             "entry_month", "exit_month")
            # 员工级列映射（身份 + 无月份的公共列），逐月复用
            emp_map: dict[str, str] = {k: v for k, v in col.items() if k in identity_keys}
            for key, header in employee_cols.items():
                emp_map.setdefault(key, header)

            skipped_blank = 0
            for row in sheet.rows:
                for month, fields in sorted(month_cols.items()):
                    # 该月整行为空（尚未入职 / 已离职）则不生成记录
                    if all(row.get(h) in (None, "") for h in fields.values()):
                        skipped_blank += 1
                        continue
                    col_map = {**emp_map, **fields}
                    rec = make_record(row, col_map)
                    rec["month"] = month
                    records.append(rec)
            if skipped_blank:
                report.notes.append(f"已跳过 {skipped_blank} 个整月为空的记录（未入职/已离职月份）")

    else:  # single
        month = ctx.month or 1
        report.notes.append(f"单月表，月份按 {month} 月处理")
        for row in sheet.rows:
            rec = base_record(row)
            rec["month"] = int(rec.get("month") or month)
            records.append(rec)

    report.records = len(records)
    return records, report
